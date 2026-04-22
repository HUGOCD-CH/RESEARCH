import io
import re
import datetime
from flask import Flask, jsonify, redirect, render_template, request, send_file, url_for
import auth
import config
import mail

app = Flask(__name__)
app.secret_key = config.SECRET_KEY
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"


def _require_auth():
    if not auth.is_ready():
        return redirect(url_for("index"))
    return None


@app.route("/")
def index():
    state = auth.get_state()
    if state["status"] == "active":
        return redirect(url_for("inbox"))
    return render_template("index.html", state=state)


# ── Auth flow ─────────────────────────────────────────────────────────────────

@app.route("/auth/start")
def auth_start():
    auth.start_login()
    return render_template("waiting.html")


@app.route("/auth/poll")
def auth_poll():
    state = auth.get_state()
    return jsonify({
        "status": state["status"],
        "error":  state.get("error"),
        "user":   state.get("user"),
    })


@app.route("/auth/logout")
def logout():
    auth.clear()
    return redirect(url_for("index"))


# ── Mail views ────────────────────────────────────────────────────────────────

@app.route("/inbox")
def inbox():
    redir = _require_auth()
    if redir:
        return redir

    page_num = max(int(request.args.get("page", 1)), 1)
    skip = (page_num - 1) * config.MESSAGES_PER_PAGE

    try:
        data = mail.get_messages(top=config.MESSAGES_PER_PAGE, skip=skip)
        user = mail.get_user_profile()
    except Exception as exc:
        return render_template("error.html", message=str(exc))

    messages = data.get("value", []) if isinstance(data, dict) else []
    return render_template(
        "inbox.html",
        messages=messages,
        page=page_num,
        has_next=len(messages) == config.MESSAGES_PER_PAGE,
        has_prev=page_num > 1,
        user=user,
    )


@app.route("/email/<message_id>")
def view_email(message_id):
    redir = _require_auth()
    if redir:
        return redir

    try:
        message = mail.get_message(message_id)
        user = mail.get_user_profile()
        mail.mark_as_read(message_id)
    except Exception as exc:
        return render_template("error.html", message=str(exc))

    body = message.get("body", {})
    return render_template(
        "email.html",
        message=message,
        body_content=body.get("content", ""),
        body_is_html=body.get("contentType", "text").lower() == "html",
        user=user,
    )


@app.route("/export/today")
def export_today():
    redir = _require_auth()
    if redir:
        return redir

    try:
        messages = mail.get_todays_messages()
    except Exception as exc:
        return render_template("error.html", message=str(exc))

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Today's Emails"

    headers = ["From Name", "From Email", "Subject", "Received At", "Read", "Body"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="0078D4")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A2"

    _tag_re = re.compile(r"<[^>]+>")

    def strip_html(html: str) -> str:
        text = _tag_re.sub(" ", html)
        text = re.sub(r"&nbsp;", " ", text)
        text = re.sub(r"&amp;", "&", text)
        text = re.sub(r"&lt;", "<", text)
        text = re.sub(r"&gt;", ">", text)
        text = re.sub(r"&quot;", '"', text)
        return re.sub(r" {2,}", " ", text).strip()

    for row_num, msg in enumerate(messages, 2):
        sender = msg.get("from", {}).get("emailAddress", {})
        body = msg.get("body", {})
        raw_body = body.get("content", "")
        plain_body = strip_html(raw_body) if body.get("contentType", "text") == "html" else raw_body

        ws.cell(row=row_num, column=1, value=sender.get("name", ""))
        ws.cell(row=row_num, column=2, value=sender.get("address", ""))
        ws.cell(row=row_num, column=3, value=msg.get("subject", ""))
        ws.cell(row=row_num, column=4, value=msg.get("receivedDateTime", "")[:19].replace("T", " "))
        ws.cell(row=row_num, column=5, value="Yes" if msg.get("isRead") else "No")
        body_cell = ws.cell(row=row_num, column=6, value=plain_body[:32767])
        body_cell.alignment = Alignment(wrap_text=True)

    col_widths = [25, 35, 50, 20, 6, 80]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    today_str = datetime.date.today().strftime("%Y-%m-%d")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"emails_{today_str}.xlsx",
    )


if __name__ == "__main__":
    app.run(debug=False, port=5000, threaded=True)
